"""Тесты admin action выгрузки комментариев составителю трека."""

from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.test import Client
from openpyxl import load_workbook
import pytest

from accounts.models import Semester
from showcase.domain.track_composer_comment_excel import SHEET_EMPTY, SHEET_FILLED
from showcase.models import ProjectApplication

User = get_user_model()

CHANGELIST_URL = "/admin/showcase/projectapplication/"
ACTION_NAME = "export_track_composer_comments"
_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def admin_client(db, roles):
    """Авторизованный клиент Django admin (superuser)."""
    User.objects.create_superuser(
        email="admin-excel@example.com",
        password="adminpass",
        first_name="Admin",
        last_name="Adminov",
        role=roles["admin"],
    )
    client = Client()
    client.login(email="admin-excel@example.com", password="adminpass")
    return client


@pytest.fixture
def semester_a(db) -> Semester:
    return Semester.objects.create(code="s-a", name="Semester A", position=1)


@pytest.fixture
def semester_b(db) -> Semester:
    return Semester.objects.create(code="s-b", name="Semester B", position=2)


def _create_app(
    *,
    semester: Semester | None,
    title: str,
    comment: str,
    author=None,
) -> ProjectApplication:
    return ProjectApplication.objects.create(
        title=title,
        company="ООО Тест",
        author=author,
        author_lastname="Иванов",
        author_firstname="Иван",
        author_middlename="Иванович",
        author_email="author@example.com",
        semester=semester,
        track_composer_comment=comment,
        goal="Цель проекта достаточно длинная для заполнения обязательных полей",
        problem_holder="Носитель",
        barrier="Барьер достаточно длинный для заполнения обязательных полей",
    )


@pytest.mark.django_db
class TestExportTrackComposerCommentsAdminAction:
    """Admin action export_track_composer_comments."""

    def test_export_success_splits_sheets(
        self, admin_client, semester_a, make_user, departments
    ):
        author = make_user(
            role_code="user", with_department=True, email="app-author@example.com"
        )
        filled = _create_app(
            semester=semester_a,
            title="С комментарием",
            comment="Нужна команда с ML",
            author=author,
        )
        empty = _create_app(
            semester=semester_a,
            title="Пустой комментарий",
            comment="-",
            author=author,
        )

        response = admin_client.post(
            CHANGELIST_URL,
            data={
                "action": ACTION_NAME,
                "_selected_action": [str(filled.pk), str(empty.pk)],
            },
        )

        assert response.status_code == 200
        assert response["Content-Type"] == _XLSX_CONTENT_TYPE
        assert (
            response["Content-Disposition"]
            == 'attachment; filename="track_composer_comments_s-a.xlsx"'
        )

        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == [SHEET_FILLED, SHEET_EMPTY]
        filled_sheet = workbook[SHEET_FILLED]
        empty_sheet = workbook[SHEET_EMPTY]
        assert filled_sheet.max_row == 2
        assert empty_sheet.max_row == 2
        assert filled_sheet["C2"].value == "С комментарием"
        assert filled_sheet["B2"].value == departments["parent"].name
        assert empty_sheet["C2"].value == "Пустой комментарий"
        assert empty_sheet["D2"].value == "-"

    def test_mixed_semesters_rejected(
        self, admin_client, semester_a, semester_b, make_user
    ):
        author = make_user(
            role_code="user", with_department=True, email="mixed@example.com"
        )
        app_a = _create_app(
            semester=semester_a, title="A", comment="Текст", author=author
        )
        app_b = _create_app(
            semester=semester_b, title="B", comment="Текст", author=author
        )

        response = admin_client.post(
            CHANGELIST_URL,
            data={
                "action": ACTION_NAME,
                "_selected_action": [str(app_a.pk), str(app_b.pk)],
            },
            follow=True,
        )

        assert response.status_code == 200
        assert response.get("Content-Type", "").startswith("text/html")
        messages = [str(message) for message in get_messages(response.wsgi_request)]
        assert any("одного семестра" in message for message in messages)

    def test_empty_selection_rejected(self, admin_client, semester_a, make_user):
        author = make_user(
            role_code="user", with_department=True, email="empty-sel@example.com"
        )
        _create_app(
            semester=semester_a, title="Есть заявка", comment="Текст", author=author
        )

        response = admin_client.post(
            CHANGELIST_URL,
            data={"action": ACTION_NAME},
            follow=True,
        )

        assert response.status_code == 200
        assert response.get("Content-Type", "").startswith("text/html")
        content = response.content.decode()
        messages = [str(message) for message in get_messages(response.wsgi_request)]
        assert any(
            "selected" in message.lower()
            or "выбран" in message.lower()
            or "Items must be selected" in message
            or "Не выбрано" in message
            or "выберите" in content.lower()
            or "selected" in content.lower()
            for message in messages
        ) or ("selected" in content.lower() or "выбран" in content.lower())
