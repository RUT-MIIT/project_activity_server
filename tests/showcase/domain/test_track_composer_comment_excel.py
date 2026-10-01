"""Тесты domain-модуля Excel-выгрузки комментариев составителю трека."""

from io import BytesIO

from openpyxl import load_workbook
import pytest

from accounts.models import Semester
from showcase.domain.track_composer_comment_excel import (
    HEADERS,
    SHEET_EMPTY,
    SHEET_FILLED,
    application_to_row,
    build_track_composer_comment_xlsx,
    is_empty_track_composer_comment,
    resolve_author_institute_name,
)
from showcase.models import ProjectApplication


@pytest.mark.parametrize(
    ("text", "expected_empty"),
    [
        ("", True),
        ("   ", True),
        ("-", True),
        ("—", True),
        ("–", True),
        (".", True),
        ("…", True),
        ("нет", True),
        ("Нет", True),
        ("н/д", True),
        ("н.д.", True),
        ("n/a", True),
        ("N/A", True),
        ("na", True),
        ("нет комментария", True),
        ("без комментария", True),
        ("ML", False),
        ("Нужна команда с ML", False),
        (None, True),
    ],
)
def test_is_empty_track_composer_comment(
    text: str | None, expected_empty: bool
) -> None:
    """Классификация пустых и заполненных комментариев по спеке."""
    assert is_empty_track_composer_comment(text) is expected_empty


@pytest.mark.django_db
class TestTrackComposerCommentExcelBuild:
    """Сборка строк и xlsx-файла."""

    @pytest.fixture
    def semester(self, db) -> Semester:
        return Semester.objects.create(code="s-excel", name="S Excel", position=1)

    def _create_app(
        self,
        *,
        semester: Semester,
        title: str,
        comment: str,
        author=None,
        last_name: str = "Иванов",
        first_name: str = "Иван",
        middle_name: str = "Иванович",
    ) -> ProjectApplication:
        return ProjectApplication.objects.create(
            title=title,
            company="ООО Тест",
            author=author,
            author_lastname=last_name,
            author_firstname=first_name,
            author_middlename=middle_name,
            author_email="author@example.com",
            semester=semester,
            track_composer_comment=comment,
            goal="Цель проекта достаточно длинная для заполнения обязательных полей",
            problem_holder="Носитель",
            barrier="Барьер достаточно длинный для заполнения обязательных полей",
        )

    def test_build_splits_filled_and_empty_sheets(
        self, semester, make_user, departments
    ):
        author = make_user(
            role_code="user", with_department=True, email="author-excel@example.com"
        )
        filled = self._create_app(
            semester=semester,
            title="С комментарием",
            comment="Нужна команда с ML",
            author=author,
        )
        empty = self._create_app(
            semester=semester,
            title="Без комментария",
            comment="-",
            author=author,
            last_name="Петров",
            first_name="Пётр",
            middle_name="",
        )

        payload = build_track_composer_comment_xlsx([filled, empty])
        workbook = load_workbook(BytesIO(payload))

        assert workbook.sheetnames == [SHEET_FILLED, SHEET_EMPTY]
        filled_sheet = workbook[SHEET_FILLED]
        empty_sheet = workbook[SHEET_EMPTY]

        assert tuple(cell.value for cell in filled_sheet[1]) == HEADERS
        assert tuple(cell.value for cell in empty_sheet[1]) == HEADERS

        filled_row = tuple(cell.value for cell in filled_sheet[2])
        assert filled_row[0] == "Иванов Иван Иванович"
        assert filled_row[1] == departments["parent"].name
        assert filled_row[2] == "С комментарием"
        assert filled_row[3] == "Нужна команда с ML"
        assert filled_sheet.max_row == 2

        empty_row = tuple(cell.value for cell in empty_sheet[2])
        assert empty_row[0] == "Петров Пётр"
        assert empty_row[1] == departments["parent"].name
        assert empty_row[2] == "Без комментария"
        assert empty_row[3] == "-"
        assert empty_sheet.max_row == 2

    def test_resolve_institute_from_root_department(
        self, semester, make_user, departments
    ):
        author = make_user(
            role_code="user", with_department=True, email="author-root@example.com"
        )
        app = self._create_app(
            semester=semester,
            title="Проект",
            comment="Текст",
            author=author,
        )

        assert author.department_id == departments["child"].id
        assert resolve_author_institute_name(app) == departments["parent"].name
        assert application_to_row(app)[1] == departments["parent"].name

    def test_resolve_institute_empty_without_author(self, semester):
        app = self._create_app(
            semester=semester,
            title="Без автора",
            comment="Текст",
            author=None,
        )

        assert resolve_author_institute_name(app) == ""
        assert application_to_row(app)[1] == ""

    def test_resolve_institute_empty_without_department(self, semester, make_user):
        author = make_user(
            role_code="user", with_department=False, email="no-dept@example.com"
        )
        app = self._create_app(
            semester=semester,
            title="Без отдела",
            comment="Текст",
            author=author,
        )

        assert author.department_id is None
        assert resolve_author_institute_name(app) == ""
